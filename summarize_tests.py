"""从真实pytest结果生成根目录摘要和补充表：python summarize_tests.py。"""
import csv
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).parent
OUT = ROOT / "测试结果"
FINAL = OUT / "final"


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def sheet(book, name, headers, rows):
    ws = book.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="245A81")
        cell.font = Font(color="FFFFFF", bold=True)
    from openpyxl.utils import get_column_letter
    for i in range(1, len(headers)+1):
        ws.column_dimensions[get_column_letter(i)].width = 24 if i < 3 else 55
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    with (OUT / (name + ".csv")).open("w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows([headers] + rows)


def main():
    final = read(FINAL / "results.json")
    rows = final["rows"]
    assert rows, "最终测试尚未完成"
    workbook = Workbook()
    workbook.remove(workbook.active)
    totals = Counter(r[6] for r in rows)
    stat_rows = []
    for dim, idx in [("方法", 7), ("目标", 8), ("阶段", 9)]:
        for name in sorted({r[idx] for r in rows}):
            count = Counter(r[6] for r in rows if r[idx] == name)
            stat_rows.append([dim, name, sum(count.values()), *[count[v] for v in ("通过", "失败", "错误", "跳过")]])
    sheet(workbook, "分类汇总", ["维度", "类别", "用例数", "通过", "失败", "错误", "跳过"], stat_rows)
    matrix = []
    for method, goal in sorted({(r[7], r[8]) for r in rows}):
        count = Counter(r[6] for r in rows if r[7:9] == [method, goal])
        matrix.append([method, goal, sum(count.values()), *[count[v] for v in ("通过", "失败", "错误", "跳过")]])
    sheet(workbook, "方法与目标矩阵", ["测试方法", "测试目标", "用例数", "通过", "失败", "错误", "跳过"], matrix)

    performance = [read(p) for p in sorted(FINAL.glob("performance-*.json"))]
    perf_rows = [[p.get("endpoint", "GET /api/v1/health"), p["concurrency"], p["requests"], p["successful"], p["TPS"], p["mean_ms"], p["P50_ms"], p["P95_ms"], p["scope"]] for p in performance]
    sheet(workbook, "性能采样", ["接口", "并发数", "请求数", "成功数", "TPS", "平均ms", "P50ms", "P95ms", "测量范围"], perf_rows)

    baseline = []
    base_rows = []
    for name in ("baseline-frontend", "baseline-java", "baseline-api"):
        doc = read(OUT / name / "results.json")
        base_rows += doc["rows"]
        count = Counter(r[6] for r in doc["rows"])
        baseline.append([name, len(doc["rows"]), *[count[v] for v in ("通过", "失败", "错误", "跳过")]])
    baseline.append(["final", len(rows), *[totals[v] for v in ("通过", "失败", "错误", "跳过")]])
    sheet(workbook, "基线与回归", ["执行批次", "用例数", "通过", "失败", "错误", "跳过"], baseline)

    final_by_id = {r[0]: r for r in rows}
    defects = []
    for row in base_rows:
        if row[6] == "失败":
            current = final_by_id.get(row[0])
            defects.append([row[0], row[1], row[2], row[3], row[4], row[5], "失败", current[6] if current else "未覆盖", current[11] if current else ""])
    sheet(workbook, "缺陷修复回归", ["编号", "测试模块", "用例名称", "输入", "预期输出", "基线实际输出", "基线结果", "回归结果", "代码定位"], defects)

    manual = [
        ["MAN-01", "微信登录", "真实wx.login凭证交换", "有效、过期、重复使用code", "有效code成功；其余明确拒绝", "未连接真实微信凭证服务，本次接口数据采用H5测试登录", "未执行"],
        ["MAN-02", "iOS兼容", "iPhone微信完整流程", "登录→上传→审核→版本→任务", "显示、交互、下载预览均正确", "当前无iOS真机执行证据", "未执行"],
        ["MAN-03", "Android兼容", "Android微信完整流程", "不同屏幕和字体大小", "布局正确且无功能异常", "当前无Android真机执行证据", "未执行"],
        ["MAN-04", "消息UI", "消息标题已读按钮与胶囊不重叠", "有未读消息；不同机型胶囊位置", "按钮完整可见可点击；点击后未读归零", "编译产物验证完成，尚未真机目视验收", "未执行"],
        ["MAN-05", "文档预览", "真实PDF/Office/图片解码", "真实合法格式及损坏文件", "支持格式可打开；损坏文件有提示", "扩展名测试样本仅为小字节数据，不能代表文档解码通过", "未执行"],
        ["MAN-06", "网络可靠性", "弱网离线重连及长时间会话", "断网、切网、多次登录过期", "提示准确；恢复后可继续操作", "本次仅单元测试网络失败分支，未做真机弱网长稳测试", "未执行"],
        ["MAN-07", "归档包", "完整归档ZIP下载", "结项项目包含多个文件版本", "规划能力：生成可解压完整包", "静态实现仅更新项目archived状态，未见真实归档压缩包生成", "当前未实现"],
        ["MAN-08", "登录刷新", "真实refreshToken刷新会话", "有效及过期refreshToken", "规划能力：返回可用新JWT", "UserController.refreshToken目前返回空Token的演示数据", "当前未实现"],
        ["MAN-09", "安全验收", "其他资源与密码分享完整权限审计", "其余读写接口、密码分享、原始/files链接", "按角色、所属小组和分享密码完整隔离", "本次只验证已列出的权限用例；其余接口及静态文件链接仍需单独审计", "未执行"],
        ["MAN-10", "前端功能", "首页审核、重新提交及成果管理入口", "点击相关页面按钮", "规划能力：完成真实业务请求", "静态代码仍存在TODO/开发中提示；后端接口测试不代表这些页面入口可用", "当前部分未实现"],
    ]
    sheet(workbook, "待人工验收", ["编号", "测试模块", "用例名称", "输入", "预期输出", "实际输出", "是否通过"], manual)

    coverage = ET.parse(FINAL / "java-coverage/jacoco.xml").getroot()
    coverage_rows = []
    for cls in coverage.iter("class"):
        line = next((c for c in cls.findall("counter") if c.get("type") == "LINE"), None)
        branch = next((c for c in cls.findall("counter") if c.get("type") == "BRANCH"), None)
        if line is None:
            continue
        covered, missed = int(line.get("covered")), int(line.get("missed"))
        if covered == 0:
            continue
        covered_b = int(branch.get("covered")) if branch is not None else 0
        missed_b = int(branch.get("missed")) if branch is not None else 0
        coverage_rows.append([cls.get("name"), covered, missed, round(covered/(covered+missed)*100, 2), covered_b, missed_b,
            round(covered_b/(covered_b+missed_b)*100, 2) if covered_b+missed_b else "不适用"])
    sheet(workbook, "Java单元覆盖率", ["类", "已覆盖行", "未覆盖行", "行覆盖率%", "已覆盖分支", "未覆盖分支", "分支覆盖率%"], coverage_rows)
    workbook.save(OUT / "测试汇总与补充表.xlsx")
    env = read(FINAL / "environment.json")
    total_line = next(c for c in coverage.findall("counter") if c.get("type") == "LINE")
    covered, missed = int(total_line.get("covered")), int(total_line.get("missed"))
    controllers = ROOT / "knife4j-spring/src/main/java/com/rocgit/knife4j/controller"
    endpoints = {p.stem: len(re.findall(r"@(Get|Post|Put|Delete|Patch)Mapping\b", p.read_text(encoding="utf-8-sig"))) for p in controllers.glob("*.java")}
    summary = {"total": len(rows), "statuses": dict(totals), "methods": dict(Counter(r[7] for r in rows)),
        "goals": dict(Counter(r[8] for r in rows)), "stages": dict(Counter(r[9] for r in rows)),
        "baseline": baseline, "baseline_failed": len(defects), "performance": performance,
        "java_unit_line_coverage_percent": round(covered/(covered+missed)*100, 2),
        "controllers": len(endpoints), "http_endpoints": sum(endpoints.values()), "endpoint_inventory": endpoints,
        "git_head": env["git_head"], "cleanup": env.get("cleanup"), "time": final["time"]}
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    import sys
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()

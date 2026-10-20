"""独立测试实例生命周期、pytest 证据采集和 Excel/CSV 报表。"""
import csv
import json
import os
from pathlib import Path
from datetime import datetime, timezone, timedelta
from collections import Counter
import pytest
from testing.support import describe

ROOT = Path(__file__).parent
ROWS = []


@pytest.fixture(scope="session")
def lab(pytestconfig):
    if not pytestconfig.getoption("--remote"):
        pytest.skip("未启用 --remote；未运行，不计为通过")
    from testing.remote_lab import RemoteLab
    key = pytestconfig.getoption("--ssh-key")
    if not key or not Path(key).is_file():
        pytest.fail("--remote 需要有效的 --ssh-key 私钥路径")
    instance = RemoteLab(pytestconfig.getoption("--ssh-host"), pytestconfig.getoption("--ssh-user"), key)
    out = ROOT / pytestconfig.getoption("--report-dir")
    out.mkdir(parents=True, exist_ok=True)
    try:
        instance.start()
        (out / "environment.json").write_text(describe(instance.evidence), encoding="utf-8")
        yield instance
    finally:
        instance.close(out)


@pytest.fixture
def api(lab, request):
    from testing.remote_lab import Api
    return Api(lab, request.node)


@pytest.fixture
def world(api, request):
    """每条业务用例创建独立用户、小组和项目，无用例执行顺序依赖。"""
    import uuid
    tag = uuid.uuid4().hex[:10]
    users = []
    for role in ("leader", "member", "outsider"):
        users.append(api.ok("POST", "/user/login", json={"code": f"t_{tag}_{role}",
            "nickname": f"测试_{role}", "avatarUrl": "https://example.invalid/a.png", "source": "h5"}))
    leader, member, outsider = users
    api.token = leader["token"]
    group = api.ok("POST", "/group/create", json={"name": "测试组_" + tag, "description": "隔离测试"})
    api.ok("POST", "/group/join", token=member["token"], json={"inviteCode": group["inviteCode"]})
    project = api.ok("POST", "/project/create", json={"groupId": group["groupId"], "name": "测试项目_" + tag,
        "stages": [{"name": "分析", "orderNum": 1}, {"name": "测试", "orderNum": 2}],
        "members": [{"userId": leader["userId"], "role": "leader"}, {"userId": member["userId"], "role": "member"}]})
    request.node.observed = []
    return dict(leader=leader, member=member, outsider=outsider, group=group, project=project,
                pid=project["projectId"], gid=group["groupId"])


@pytest.hookimpl(hookwrapper=True)
def pytest_runtest_makereport(item, call):
    report = (yield).get_result()
    if report.when == "setup" and report.passed:
        return
    if report.when == "teardown":
        if report.failed:
            ROWS.append(["ENV-CLEANUP", "环境", "测试资源清理", "会话结束", "清理成功", str(report.longrepr),
                         "错误", "灰盒", "可靠性", "环境", report.duration, item.nodeid])
        return
    marker = item.get_closest_marker("case")
    m = marker.kwargs if marker else {}
    status = "通过" if report.passed else "跳过" if report.skipped else "错误" if report.when == "setup" else "失败"
    observed = getattr(item, "observed", [])
    actual = describe(observed) if observed else "断言全部满足" if report.passed else ""
    if not report.passed:
        actual += "\n" + str(report.longrepr)[-7000:]
    case_id = m.get("id", item.name)
    if hasattr(item, "callspec") and item.callspec.id != case_id:
        case_id += "[" + item.callspec.id + "]"
    ROWS.append([case_id, m.get("module", "未分类"), m.get("name", item.name),
        m.get("inputs", "见代码"), m.get("expected", "断言满足"), actual, status,
        m.get("method", "白盒"), m.get("goal", "功能"), m.get("stage", "单元测试"),
        round(report.duration, 4), item.nodeid])


def pytest_sessionfinish(session, exitstatus):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    out = ROOT / session.config.getoption("--report-dir")
    out.mkdir(parents=True, exist_ok=True)
    headers = ["编号", "测试模块", "用例名称", "输入", "预期输出", "实际输出", "是否通过", "测试方法", "测试目标", "测试阶段", "耗时秒", "代码定位"]
    workbook = Workbook()
    workbook.remove(workbook.active)
    groups = {"全部用例": ROWS, **{m + "测试": [r for r in ROWS if r[7] == m] for m in ("黑盒", "白盒", "灰盒")}}
    for name, rows in groups.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append([ILLEGAL_CHARACTERS_RE.sub("", str(x)) if isinstance(x, str) else x for x in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="245A81")
        for col, width in zip("ABCDEFGHIJKL", [20, 18, 40, 45, 55, 85, 12, 14, 16, 16, 14, 65]):
            sheet.column_dimensions[col].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            row[6].font = Font(color={"通过": "16803C", "失败": "CE2020", "错误": "CE2020"}.get(row[6].value, "777777"))
        with (out / (name + ".csv")).open("w", newline="", encoding="utf-8-sig") as stream:
            csv.writer(stream).writerows([headers] + [[("'" + x) if isinstance(x, str) and x.startswith(("=", "+", "-", "@")) else x for x in row] for row in rows])
    stats = workbook.create_sheet("分类统计")
    stats.append(["维度", "类别", "通过", "失败", "错误", "跳过", "总数"])
    for dim, idx in (("测试方法", 7), ("测试目标", 8), ("测试阶段", 9)):
        for label in sorted({r[idx] for r in ROWS}):
            counts = Counter(r[6] for r in ROWS if r[idx] == label)
            stats.append([dim, label, *[counts[s] for s in ("通过", "失败", "错误", "跳过")], sum(counts.values())])
    workbook.save(out / "测试结果.xlsx")
    stamp = datetime.now(timezone(timedelta(hours=8))).isoformat()
    (out / "results.json").write_text(json.dumps({"time": stamp, "exitstatus": int(exitstatus),
        "headers": headers, "rows": ROWS}, ensure_ascii=False, indent=2), encoding="utf-8")
